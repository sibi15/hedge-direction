"""Score the filled-in hand-rating sheet.

Usage: python score_ratings.py hand_rating_sheet.xlsx hand_rating_key.csv
Prints rater agreement (Cohen's kappa), keyword-judge accuracy against the agreed labels,
and the human hedge/direct rates per condition (to compare with the judge's rates in the paper).
"""
import csv
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

sheet, keyf = sys.argv[1], sys.argv[2]
key = {r["id"]: r for r in csv.DictReader(open(keyf, encoding="utf-8"))}
wb = load_workbook(sheet)
CODE = {"d": "direct", "h": "hedge", "x": "disclaimer"}


def read(name):
    out = {}
    for row in wb[name].iter_rows(min_row=2, values_only=True):
        if row[0] and row[3]:
            out[row[0]] = CODE.get(str(row[3]).strip().lower())
    return out


A, B = read("Rater A"), read("Rater B")
both = [i for i in key if A.get(i) and B.get(i)]


def kappa(x, y):
    n = len(x)
    po = sum(a == b for a, b in zip(x, y)) / n
    cx, cy = Counter(x), Counter(y)
    pe = sum(cx[c] * cy[c] for c in set(x) | set(y)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


a = [A[i] for i in both]
b = [B[i] for i in both]
print(f"rated by both: {len(both)} / {len(key)}")
print(f"rater agreement: {sum(x == y for x, y in zip(a, b)) / len(both):.0%}, Cohen's kappa {kappa(a, b):.2f}")

agreed = [i for i in both if A[i] == B[i]]
gold = [A[i] for i in agreed]
judge = [key[i]["judge_label"] for i in agreed]
print(f"agreed rows: {len(agreed)}")
print(f"keyword judge vs agreed human label: accuracy {sum(g == j for g, j in zip(gold, judge)) / len(agreed):.0%}, "
      f"kappa {kappa(gold, judge):.2f}")
conf = Counter(zip(gold, judge))
print("confusion (human -> judge):", dict(conf))

by = defaultdict(lambda: [Counter(), Counter()])
for i in agreed:
    c = key[i]["condition"]
    by[c][0][A[i]] += 1
    by[c][1][key[i]["judge_label"]] += 1
print("\nper condition (agreed rows): human counts | judge counts")
for c, (h, j) in sorted(by.items()):
    print(f"  {c:14s} human {dict(h)} | judge {dict(j)}")
