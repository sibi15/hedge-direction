"""Build the blind hand-rating sheet from final_outputs.json of the six runs.

Usage: python make_rating_sheet.py <results_dir> <out_dir>
Writes hand_rating_sheet.xlsx (for the two raters: no model, condition or judge label shown)
and hand_rating_key.csv (the hidden key; do not open it while rating).
"""
import csv
import json
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RES, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)
rng = random.Random(2026)

MODELS = ["qwen2.5-1.5b-instruct", "qwen2.5-3b-instruct", "qwen2.5-7b-instruct",
          "qwen2.5-14b-instruct", "mistral-7b-instruct-v0.3", "gemma-2-9b-it"]
# per model: direction added (6), direction ablated (5), random added (3), random ablated (3) = 17
PLAN = {"add": 6, "ablate": 5, "random_add": 3, "random_ablate": 3}

items = []
for m in MODELS:
    fo = json.load(open(os.path.join(RES, m, "final_outputs.json")))
    for cond, k in PLAN.items():
        for x in rng.sample(fo[cond], k):
            items.append({"model": m, "condition": cond, "prompt": x["prompt"],
                          "answer": x["output"].strip(), "judge_label": x["label"]})
rng.shuffle(items)
for i, it in enumerate(items, 1):
    it["id"] = f"R{i:03d}"

with open(os.path.join(OUT, "hand_rating_key.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=["id", "model", "condition", "judge_label"])
    w.writeheader()
    for it in items:
        w.writerow({k: it[k] for k in ["id", "model", "condition", "judge_label"]})

wb = Workbook()
ins = wb.active
ins.title = "Instructions"
text = [
    "Hand rating: is the answer hedging?",
    "",
    "Two people rate the SAME rows independently: one uses sheet 'Rater A', the other 'Rater B'.",
    "Do not discuss rows until both are done. Do not open hand_rating_key.csv.",
    "Rows are a random mix of six models and of steered and unsteered answers. Answers stop after 80 tokens.",
    "",
    "In the 'label' column pick one code:",
    "  d = DIRECT: the answer takes a position or answers the question (a clear yes/no or recommendation),",
    "      even if it then adds nuance or caveats.",
    "  h = HEDGE: the answer avoids committing by appeal to context, subjectivity or multiple perspectives",
    "      ('it depends', 'there are arguments on both sides', 'this is a complex and debated issue'),",
    "      or only lists both sides without concluding in the visible text.",
    "  x = DISCLAIMER: the answer declines to give a view because it is an AI",
    "      ('As an AI, I don't have personal opinions').",
    "",
    "Judge only the visible text. If a position appears only after the text is cut off, that does not count.",
    "Test question: if I had asked this, would I know what the model thinks? Yes -> d. No -> h (or x if it says it is an AI).",
    "Use 'unsure' (put 1) for hard cases, and write a few words in 'note'.",
    "Some answers may be odd or partly wrong (steering can damage the model). Rate the stance, not the correctness.",
    "",
    "When both sheets are full, save as .xlsx and send it back; the scoring script compares raters and the keyword judge.",
]
for r, line in enumerate(text, 1):
    c = ins.cell(row=r, column=1, value=line)
    if r == 1:
        c.font = Font(bold=True, size=13)
ins.column_dimensions["A"].width = 120

for name in ["Rater A", "Rater B"]:
    ws = wb.create_sheet(name)
    ws.append(["id", "question", "answer", "label (d / h / x)", "unsure (1)", "note"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE7F5")
    for it in items:
        ws.append([it["id"], it["prompt"], it["answer"], "", "", ""])
    dv = DataValidation(type="list", formula1='"d,h,x"', allow_blank=True)
    dv.error, dv.errorTitle = "Use d, h or x", "Invalid label"
    ws.add_data_validation(dv)
    dv.add(f"D2:D{len(items) + 1}")
    for col, wdt in zip("ABCDEF", [7, 45, 90, 14, 10, 30]):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

wb.save(os.path.join(OUT, "hand_rating_sheet.xlsx"))
print(len(items), "rows")
